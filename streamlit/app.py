"""Masar — Saudi Industrial Mobility Intelligence dashboard (SDA-21).

Answers the Operations Manager's questions:
  1. Which industrial cities and regions are busiest, and which grew fastest Q1 → Q2?
  2. How does movement change month by month?
  3. What was the weather and event context? (context, not cause)

Run:   streamlit run app.py
Data:  see data_source.py  (MASAR_DATA_SOURCE = auto | bigquery | local | snapshot)
"""
from __future__ import annotations

import os
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

import data_source
import metrics
import theme as T

st.set_page_config(page_title="Masar · Industrial Mobility", page_icon=T.FAVICON, layout="wide")
st.markdown(T.CSS, unsafe_allow_html=True)

PROJECT_DIR = Path(os.environ.get("MASAR_PROJECT_DIR", Path(__file__).resolve().parent.parent))
SOURCE_SHORT = {"bigquery": "Live BigQuery", "local": "Local pipeline files", "snapshot": "Bundled snapshot"}

# Readable names for every column a viewer can see in a table.
COLUMN_LABELS = {
    "city_code": "City code", "city_name": "Industrial city", "region_code": "Region code",
    "region_name": "Region", "month_start": "Month", "quarter": "Quarter", "entries": "Entries",
    "exits": "Exits", "total_movements": "Total movements", "entry_exit_ratio": "Entry / exit ratio",
    "in_both_quarters": "In both quarters", "quality_flag": "Quality flag", "avg_temp_c": "Avg temp (°C)",
    "max_temp_c": "Max temp (°C)", "hot_days": "Hot days", "precipitation_mm": "Rain (mm)",
    "event_count": "Event links", "reason": "Why flagged", "missing_months": "Missing months",
    "months_with_data": "Months with data", "file_name": "File", "sheet_name": "Worksheet",
}
INT_COLS = {"entries", "exits", "total_movements", "event_count", "hot_days", "months_with_data"}
DEC_COLS = {"entry_exit_ratio": "{:.2f}", "avg_temp_c": "{:.1f}", "max_temp_c": "{:.1f}", "precipitation_mm": "{:.1f}"}


def num(n) -> str:
    if n is None or pd.isna(n):
        return "—"
    n = float(n)
    if abs(n) >= 1e6:
        return f"{n / 1e6:.2f}M"
    if abs(n) >= 1e3:
        return f"{n / 1e3:.0f}K"
    return f"{n:,.0f}"


def short_city(name: str | None) -> str:
    if not name:
        return "—"
    for suffix in (" City for Industry and Business", " Industrial City"):
        name = name.replace(suffix, "")
    return name


def pretty_table(df: pd.DataFrame):
    """Readable headers, months as 'Jan 2026', thousands separators — still sortable."""
    out = df.copy()
    if "month_start" in out:
        out["month_start"] = pd.to_datetime(out["month_start"]).dt.strftime("%b %Y")
    fmt = {COLUMN_LABELS.get(c, c): "{:,.0f}" for c in out.columns if c in INT_COLS}
    fmt.update({COLUMN_LABELS.get(c, c): f for c, f in DEC_COLS.items() if c in out.columns})
    out = out.rename(columns=lambda c: COLUMN_LABELS.get(c, c.replace("_", " ").capitalize()))
    return out.style.format(fmt, na_rep="—")


def label_positions(values: dict, min_gap: float) -> dict:
    """Spread end-of-line labels so close values don't overlap (returns y per label)."""
    placed, last = {}, None
    for name, y in sorted(values.items(), key=lambda kv: kv[1]):
        y = y if last is None or y - last >= min_gap else last + min_gap
        placed[name], last = y, y
    return placed


def show(fig, height, title=None, **layout):
    fig = T.style(fig, height, title)
    if layout:
        fig.update_layout(**layout)
    # no width/height kwargs: older Streamlit versions treat them as deprecated Plotly kwargs and show a warning
    st.plotly_chart(fig, config={"displayModeBar": False})


def html(s: str):
    st.markdown(s, unsafe_allow_html=True)


# --------------------------------------------------------------------------- data
@st.cache_data(ttl=600, show_spinner="Loading pipeline data…")
def get_data():
    return data_source.load()


@st.cache_data(ttl=3600)
def source_inventory(project_dir: str) -> tuple[int, int | None]:
    """Excel files in the repo root and how many worksheets they hold (None if unreadable)."""
    files = sorted(Path(project_dir).glob("*.xlsx"))
    inventory = Path(project_dir) / "data" / "processed" / "workbook_inventory.csv"
    if inventory.exists():
        return len(files), len(pd.read_csv(inventory))
    try:
        from openpyxl import load_workbook
        sheets = 0
        for f in files:
            wb = load_workbook(f, read_only=True)
            sheets += len(wb.sheetnames)
            wb.close()
        return len(files), sheets if files else None
    except Exception:  # openpyxl missing or a file is locked — just hide the number
        return len(files), None


result = get_data()

if result.data is None:
    html(T.header("No data source", False, "—",
                  [("—", "—", "", False)] * 5, "Data is not available right now"))
    html(T.note("Data is not available right now.",
                f"{T.esc(result.error)}<br><b>BigQuery:</b> check your Google credentials and that dbt built "
                f"<code>{T.esc(data_source.MART_TABLE)}</code>. <b>Local:</b> run <code>python -m ingestion.run</code> "
                "from the repo root, then set <code>MASAR_DATA_SOURCE=local</code>.", "warn"))
    st.stop()

data = result.data.copy()
data["month_start"] = pd.to_datetime(data["month_start"])
live = result.source == "bigquery"

# --------------------------------------------------------------------------- filters
FILTER_KEYS = ("f_region", "f_city", "f_month")


def reset_filters():
    for key in FILTER_KEYS:
        st.session_state[key] = []


with st.sidebar:
    html(T.side_brand())
    st.markdown("### Filters")
    regs = data[["region_code", "region_name"]].drop_duplicates().sort_values("region_name")
    reg_names = dict(zip(regs.region_code, regs.region_name))
    sel_regions = st.multiselect("Region", list(reg_names), format_func=reg_names.get,
                                 placeholder="All regions", key="f_region")
    pool = data if not sel_regions else data[data.region_code.isin(sel_regions)]
    cts = pool[["city_code", "city_name"]].drop_duplicates().sort_values("city_name")
    city_names = dict(zip(cts.city_code, cts.city_name))
    # a city picked earlier may drop out when the region changes
    st.session_state["f_city"] = [c for c in st.session_state.get("f_city", []) if c in city_names]
    sel_cities = st.multiselect("Industrial city", list(city_names), format_func=city_names.get,
                                placeholder="All cities", key="f_city")
    months = sorted(data.month_start.unique())
    sel_months = st.multiselect("Month", months, format_func=lambda m: pd.Timestamp(m).strftime("%b %Y"),
                                placeholder="All months", key="f_month")
    st.button("↺  Reset filters", on_click=reset_filters, width="stretch",
              disabled=not (sel_regions or sel_cities or sel_months))
    st.caption("Empty = all. Traffic data is monthly, so there is no day-level filter.")
    st.divider()
    st.caption("Weather and events are shown as **context only** — six months of data "
               "cannot prove they cause changes in movement.")

view = metrics.filter_data(data, sel_regions, sel_cities, sel_months)
view_all_months = metrics.filter_data(data, sel_regions, sel_cities, None)

# --------------------------------------------------------------------------- header + KPIs
scope = view if not view.empty else data
avail = metrics.context_available(scope)
k = metrics.kpis(scope)
n_regions = scope.region_code.nunique()
month_span = (f"{scope.month_start.min():%b}–{scope.month_start.max():%b %Y}"
              if scope.month_start.nunique() > 1 else f"{scope.month_start.min():%b %Y}")
# Month-over-month change: only when exactly one month is picked, and like-for-like
# (cities with data in both months), because missing months are never filled in.
mc = metrics.month_change(view_all_months, sel_months[0]) if len(sel_months) == 1 else None


def pct_delta(v, tip):
    if v is None or mc is None:
        return None
    return (f"{'▲' if v >= 0 else '▼'} {abs(v):.1f}% vs {mc['prev']:%b}", "up" if v >= 0 else "down", tip)


like = (f"Like-for-like: {mc['cities']} of {mc['cities_now']} cities have data in both months"
        if mc else "")
first_month = len(sel_months) == 1 and mc is None
kpi_cards = [
    ("Total movements", num(k["total"]),
     f"like-for-like {mc['cities']}/{mc['cities_now']} cities" if mc
     else f"{num(k['entries'])} entries · {num(k['exits'])} exits", True,
     pct_delta(mc["total_pct"], like) if mc else (("no previous month", "flat", "") if first_month else None)),
    ("Busiest city", short_city(k["busiest_city"]),
     (f"{k['busiest_city_share']:.0%} of scope" if mc else f"{k['busiest_city_share']:.0%} of movements in scope")
     if k["busiest_city_share"] else "", False,
     pct_delta(mc["by_city"].get(k["busiest_city"]), "This city vs the month before") if mc else None),
    ("Busiest region", k["busiest_region"] or "—",
     (f"{k['busiest_region_share']:.0%} of scope" if mc else f"{k['busiest_region_share']:.0%} of movements in scope")
     if k["busiest_region_share"] else "", False,
     pct_delta(mc["by_region"].get(k["busiest_region"]), like) if mc else None),
    ("Entry / exit ratio", "—" if k["ratio"] is None else f"{k['ratio']:.2f}", "Balanced ≈ 1.00", False,
     (f"Δ {mc['ratio_delta']:+.2f}", "flat", like) if mc else None),
    ("Avg temperature", f"{scope.avg_temp_c.mean():.1f} °C", "Open-Meteo · context only", False,
     (f"{mc['temp_delta']:+.1f} °C", "flat", "Context only") if mc and mc["temp_delta"] is not None else None)
    if avail["weather"] else ("Cities in scope", str(k["cities"]), f"Across {n_regions} region{'s' if n_regions != 1 else ''}", False),
]
html(T.header(SOURCE_SHORT.get(result.source, result.label), live, result.updated or "unknown", kpi_cards,
              f"Vehicle entries & exits across {k['cities']} industrial "
              f"{'city' if k['cities'] == 1 else 'cities'} · {month_span}"))

if not live:
    html(T.note(f"BigQuery is not reachable right now — showing {SOURCE_SHORT.get(result.source, result.label).lower()}.",
                "Numbers come from the same pipeline output, but they may be older than the live mart. "
                "Connect Google credentials (<code>gcloud auth application-default login</code>) to read BigQuery.",
                "warn"))

if view.empty:
    html(T.note("No data for this combination of filters.", "Try removing a filter, or use <b>Reset filters</b> "
                "in the sidebar.", "warn"))
    st.stop()

# --------------------------------------------------------------------------- city profile + key findings
if len(sel_cities) == 1:  # one city: its profile replaces the cross-city findings
    html(T.city_profile(metrics.city_profile(view, view_all_months, sel_cities[0]), num))
else:
    html(T.insights(metrics.insights(view, view_all_months),
                    "From the current filters · growth and flags use all six months · context, not cause"))

# --------------------------------------------------------------------------- map + trend
left, right = st.columns([7, 5])
with left:
    with st.container(border=True):
        colour_by = st.radio("Colour cities by", ["Region", "Growth Q1 → Q2"], horizontal=True, key="map_colour",
                             label_visibility="collapsed")
        pts = metrics.map_points(view)
        pts["total_movements"] = pts["total_movements"].astype(float)
        if colour_by == "Region":
            fig = px.scatter_geo(
                pts, lat="latitude", lon="longitude", size="total_movements", color="region_name",
                color_discrete_map=T.REGION_COLORS, hover_name="city_name", size_max=40,
                hover_data={"total_movements": ":,.0f", "region_name": True,
                            "latitude": False, "longitude": False},
            )
            fig.update_traces(marker=dict(line=dict(width=1.2, color=T.BG), opacity=0.92))
        else:
            growth_map, _ = metrics.growth_q1_q2(view_all_months)
            pts = pts.merge(growth_map[["city_code", "growth_pct"]], on="city_code", how="left")
            pts["size"] = 9 + 31 * np.sqrt(pts.total_movements / pts.total_movements.max())
            lim = max(5.0, float(pts.growth_pct.abs().max())) if pts.growth_pct.notna().any() else 5.0
            fig = go.Figure()
            cmp_ = pts[pts.growth_pct.notna()]
            fig.add_trace(go.Scattergeo(
                lat=cmp_.latitude, lon=cmp_.longitude, text=cmp_.city_name, name="Compared",
                customdata=np.stack([cmp_.growth_pct, cmp_.total_movements], axis=-1) if len(cmp_) else None,
                marker=dict(size=cmp_["size"], color=cmp_.growth_pct, cmin=-lim, cmax=lim, opacity=0.95,
                            colorscale=[[0, T.DOWN], [0.5, "#c9d3e6"], [1, T.UP]],
                            line=dict(width=1.2, color=T.BG),
                            colorbar=dict(title=dict(text="Growth", font=dict(color=T.TXT, size=12)),
                                          ticksuffix="%", tickfont=dict(color=T.TXT, size=11),
                                          len=0.55, thickness=10, x=0.985, xanchor="right", y=0.04,
                                          yanchor="bottom", bgcolor="rgba(10,20,48,0.78)",
                                          bordercolor=T.LINE, borderwidth=1, xpad=10, ypad=8)),
                hovertemplate="<b>%{text}</b><br>Growth Q1 → Q2: %{customdata[0]:+.1f}%<br>"
                              "Movements: %{customdata[1]:,.0f}<extra></extra>", showlegend=False))
            rest = pts[pts.growth_pct.isna()]
            if len(rest):
                fig.add_trace(go.Scattergeo(
                    lat=rest.latitude, lon=rest.longitude, text=rest.city_name,
                    name="Not compared (one quarter only)",
                    marker=dict(size=rest["size"], color="rgba(0,0,0,0)", line=dict(width=1.6, color=T.MUTED)),
                    hovertemplate="<b>%{text}</b><br>Not compared — data in one quarter only<extra></extra>"))
        fig.update_geos(  # drawn locally: no map tiles, no API key needed
            projection_type="mercator", lataxis_range=[15.6, 32.8], lonaxis_range=[25, 65],
            center=dict(lat=24.2, lon=45),
            showland=True, landcolor=T.LAND, showocean=True, oceancolor=T.SEA,
            showlakes=False, showcountries=True, countrycolor=T.BORDER, showcoastlines=True,
            coastlinecolor=T.BORDER, showframe=False, bgcolor="rgba(0,0,0,0)", resolution=50,
            domain=dict(x=[0, 1], y=[0, 1]),
        )
        show(fig, 452, "Industrial network pulse — marker size = vehicle movements"
             + ("" if colour_by == "Region" else
                f"<br><span style='font-size:12px;color:{T.MUTED}'>Colour = growth Q1 → Q2 in average monthly "
                "movement · hollow = not compared</span>"),
             margin=dict(l=0, r=0, t=40 if colour_by == "Region" else 62, b=0),
             legend=dict(orientation="v", x=0.015, xanchor="left", y=0.03, yanchor="bottom",
                         itemsizing="constant", title=None, font=dict(color=T.TXT, size=12),
                         bgcolor="rgba(10,20,48,0.78)", bordercolor=T.LINE, borderwidth=1))
with right:
    with st.container(border=True):
        trend = metrics.monthly_trend(view, "region")
        fig = px.line(trend, x="month_start", y="total_movements", color="region_name",
                      color_discrete_map=T.REGION_COLORS, markers=True)
        fig.update_traces(line_width=2.2, marker_size=8,
                          hovertemplate="%{fullData.name}: %{y:,.0f}<extra></extra>")
        fig.update_layout(hovermode="x unified", showlegend=False)
        fig.update_xaxes(dtick="M1", tickformat="%b")
        ends = {n: g.iloc[-1] for n, g in trend.groupby("region_name")}
        ys = label_positions({n: float(r.total_movements) for n, r in ends.items()},
                             min_gap=0.06 * float(trend.total_movements.max()))
        for name, r in ends.items():  # direct labels: colour is never the only cue
            fig.add_annotation(x=r.month_start, y=ys[name], text=name, showarrow=False,
                               xanchor="left", xshift=8, font=dict(color=T.TXT, size=11))
        show(fig, 512, "Monthly movement trend by region"
                       f"<br><span style='font-size:12px;color:{T.MUTED}'>Gaps = no data reported "
                       "for that month (see Data quality)</span>", margin=dict(t=64, r=70))

# --------------------------------------------------------------------------- comparison + context
left, right = st.columns([7, 5])
with left:
    with st.container(border=True):
        log_x = st.toggle("Log scale", key="city_log",
                          help="Each step on the axis is 10× more movements, so small cities stay visible. "
                               "Bar lengths are no longer proportional — read the labels.")
        by_city = metrics.totals_by(view, "city").sort_values("total_movements")
        fig = go.Figure(go.Bar(
            x=by_city.total_movements.astype(float), y=by_city.city_name, orientation="h",
            marker=dict(color=T.ACCENT, cornerradius=4),
            text=[num(v) for v in by_city.total_movements], textposition="outside", cliponaxis=False,
            textfont=dict(color=T.TXT),
            customdata=np.stack([by_city.region_name, by_city.entries.astype(float),
                                 by_city.exits.astype(float)], axis=-1),
            hovertemplate="<b>%{y}</b> · %{customdata[0]}<br>Total: %{x:,.0f}<br>"
                          "Entries: %{customdata[1]:,.0f} · Exits: %{customdata[2]:,.0f}<extra></extra>",
        ))
        vmax, vmin = float(by_city.total_movements.max()), float(by_city.total_movements.clip(lower=1).min())
        if log_x:
            fig.update_xaxes(type="log", showticklabels=False, showgrid=False,
                             range=[np.log10(vmin) - 0.6, np.log10(vmax) + 0.45])
        else:
            fig.update_xaxes(showticklabels=False, showgrid=False, range=[0, vmax * 1.18])
        fig.update_yaxes(tickfont=dict(color=T.TXT))
        show(fig, 70 + 32 * len(by_city), "Industrial city comparison — total movements"
             + (" (log scale)" if log_x else ""))

with right:
    with st.container(border=True):
        html(T.section("Weather context", "Open-Meteo daily weather, averaged per city and month"))
        if avail["weather"]:
            ctx = metrics.monthly_context(view)
            hottest = ctx.loc[ctx.avg_temp_c.idxmax()]
            html(T.ctx_row("☀", f"{view.avg_temp_c.mean():.1f} °C average", "Across selected cities and months")
                 + T.ctx_row("↗", f"{hottest.month_start:%B} is the hottest month", f"{hottest.avg_temp_c:.1f} °C average")
                 + T.ctx_row("▲", f"{int(view.hot_days.fillna(0).sum())} hot city-days",
                             f"Days with max ≥ {metrics.HOT_DAY_C:.0f} °C"))
        else:
            html(T.ctx_row("…", "Pending", "Weather appears once weather_daily is loaded"))
    with st.container(border=True):
        html(T.section("Government events", "Public-sector event approvals linked to industrial cities"))
        if avail["events"]:
            ev = view.groupby("month_start", as_index=False)["event_count"].sum()
            top = ev.loc[ev.event_count.idxmax()]
            busiest_ev = view.groupby("city_name")["event_count"].sum().idxmax()
            html(T.ctx_row("●", f"{top.month_start:%B} has the most events",
                           f"{int(top.event_count):,} event links to selected cities")
                 + T.ctx_row("◆", short_city(busiest_ev), "Industrial city with the most linked events")
                 + T.ctx_row("i", "Linked by the city rule",
                             "Same city → that industrial city; same region → all cities in the region"))
        else:
            html(T.ctx_row("…", "Pending", "Events appear once event_days is loaded"))
    html(T.note("Context only", "Weather and events are not a proven cause of movement changes."))

# --------------------------------------------------------------------------- detail tabs
t_growth, t_region, t_quality, t_context, t_table, t_pipe = st.tabs(
    ["Growth Q1 → Q2", "Regions & entries vs exits", "Data quality", "Weather & events by month",
     "Data table", "Pipeline & sources"])

with t_growth:
    html(T.section("Which cities grew fastest from Q1 to Q2 2026?",
                   "Compares the average monthly movement per quarter, because some cities miss months. "
                   "Only cities with data in both quarters are compared."))
    growth, excluded = metrics.growth_q1_q2(view_all_months)
    if growth.empty:
        html(T.note("None of the selected cities has data in both quarters."))
    else:
        g = growth.sort_values("growth_pct")
        fig = go.Figure(go.Bar(
            x=g.growth_pct, y=g.city_name, orientation="h",
            marker=dict(color=[T.UP if v >= 0 else T.DOWN for v in g.growth_pct], cornerradius=4),
            text=[f"{v:+.1f}%" for v in g.growth_pct], textposition="outside", cliponaxis=False,
            textfont=dict(color=T.TXT),
            customdata=np.stack([g.q1_avg_monthly, g.q2_avg_monthly, g.q1_months, g.q2_months], axis=-1),
            hovertemplate=("<b>%{y}</b><br>Growth: %{x:+.1f}%<br>"
                           "Q1 avg/month: %{customdata[0]:,.0f} (%{customdata[2]} months)<br>"
                           "Q2 avg/month: %{customdata[1]:,.0f} (%{customdata[3]} months)<extra></extra>"),
        ))
        fig.add_vline(x=0, line_color=T.MUTED, line_width=1)
        fig.update_xaxes(ticksuffix="%")
        fig.update_yaxes(tickfont=dict(color=T.TXT))
        show(fig, 70 + 36 * len(g), "Change in average monthly movement, Q1 → Q2")
        short = growth[(growth.q1_months < 3) | (growth.q2_months < 3)]
        if not short.empty:
            html(T.note("Read with care: fewer than 3 months in a quarter",
                        T.esc(", ".join(short.city_name)), "warn"))
    if not excluded.empty:
        st.caption("Not compared (one quarter only): "
                   + ", ".join(f"{r.city_name} ({r.reason.lower()})" for r in excluded.itertuples()))

with t_region:
    a, b = st.columns(2)
    with a:
        by_region = metrics.totals_by(view, "region").sort_values("total_movements")
        fig = go.Figure(go.Bar(
            x=by_region.total_movements.astype(float), y=by_region.region_name, orientation="h",
            marker=dict(color=[T.REGION_COLORS.get(r, T.ACCENT) for r in by_region.region_name],
                        cornerradius=4),
            text=[num(v) for v in by_region.total_movements], textposition="outside", cliponaxis=False,
            textfont=dict(color=T.TXT), hovertemplate="%{y}: %{x:,.0f}<extra></extra>"))
        fig.update_xaxes(showticklabels=False, showgrid=False,
                         range=[0, float(by_region.total_movements.max()) * 1.2])
        fig.update_yaxes(tickfont=dict(color=T.TXT))
        show(fig, 300, "Total movements by region")
    with b:
        city = st.selectbox("City", sorted(view.city_name.unique()), key="ee_city")
        one = view[view.city_name == city].sort_values("month_start")
        fig = go.Figure()
        fig.add_trace(go.Scatter(x=one.month_start, y=one.entries.astype(float), name="Entries",
                                 mode="lines+markers", line=dict(color=T.ENTRIES, width=2.2), marker_size=8))
        fig.add_trace(go.Scatter(x=one.month_start, y=one.exits.astype(float), name="Exits",
                                 mode="lines+markers", line=dict(color=T.EXITS, width=2.2), marker_size=8))
        fig.update_traces(hovertemplate="%{fullData.name}: %{y:,.0f}<extra></extra>")
        fig.update_layout(hovermode="x unified")
        fig.update_xaxes(dtick="M1", tickformat="%b")
        show(fig, 300, f"{short_city(city)} — entries vs exits",
             legend=dict(orientation="h", x=1, xanchor="right", y=1.12, yanchor="bottom"))

with t_quality:
    html(T.section("Entry / exit ratio — a data-quality indicator first",
                   f"A city-month is flagged when entries and exits differ by more than "
                   f"{metrics.IMBALANCE_LIMIT:.0%} (ratio below {metrics.RATIO_LOW:.2f} or above "
                   f"{metrics.RATIO_HIGH:.2f}) — the same rule as the ingestion quality flag — "
                   "or when the pipeline flags repeated values."))
    heat = view.pivot_table(index="city_name", columns="month_start", values="entry_exit_ratio")
    heat.columns = [pd.Timestamp(c).strftime("%b") for c in heat.columns]
    fig = go.Figure(go.Heatmap(
        z=np.log2(heat.astype(float)).values, x=list(heat.columns), y=list(heat.index),
        text=[["" if pd.isna(v) else f"{v:.2f}" for v in row] for row in heat.values],
        texttemplate="%{text}", textfont=dict(color=T.TXT), customdata=heat.values,
        colorscale=[[0, T.DOWN], [0.5, T.NEUTRAL_MID], [1, T.UP]], zmid=0, zmin=-1.5, zmax=1.5,
        xgap=2, ygap=2,
        hovertemplate="%{y} · %{x}<br>Entry/exit ratio: %{customdata:.2f}<extra></extra>",
        colorbar=dict(title="ratio", tickvals=[-1, 0, 1], ticktext=["0.5", "1.0", "2.0"],
                      tickfont=dict(color=T.MUTED))))
    fig.update_yaxes(tickfont=dict(color=T.TXT))
    show(fig, 80 + 30 * len(heat), "Entry / exit ratio by city and month (blank = no data)")
    flags = metrics.quality_flags(view)
    html(T.section(f"Flagged city-months: {len(flags)}"))
    if not flags.empty:
        st.dataframe(pretty_table(flags), width="stretch", hide_index=True)
    gaps = metrics.missing_months(view_all_months)
    html(T.section(f"Cities with missing months: {len(gaps)}", "Never filled in — shown as gaps in the charts"))
    if not gaps.empty:
        st.dataframe(pretty_table(gaps.drop(columns="city_code")), width="stretch", hide_index=True)

with t_context:
    html(T.note("Context only", "With six months of data, weather and events cannot be shown to "
                "<b>cause</b> changes in vehicle movement.", "warn"))
    ctx = metrics.monthly_context(view)
    fig = go.Figure(go.Bar(x=ctx.month_start, y=ctx.total_movements.astype(float),
                           marker=dict(color=T.ACCENT, cornerradius=4),
                           hovertemplate="%{x|%b}: %{y:,.0f}<extra></extra>"))
    fig.update_xaxes(dtick="M1", tickformat="%b")
    show(fig, 260, "Vehicle movements per month")
    a, b = st.columns(2)
    with a:
        if avail["weather"]:
            fig = go.Figure(go.Scatter(x=ctx.month_start, y=ctx.avg_temp_c, mode="lines+markers",
                                       line=dict(color=T.EXITS, width=2.2), marker_size=8,
                                       hovertemplate="%{x|%b}: %{y:.1f} °C<extra></extra>"))
            fig.update_xaxes(dtick="M1", tickformat="%b")
            show(fig, 240, "Average temperature (°C)")
        else:
            html(T.note("Weather: pending", "Appears once weather_daily is loaded."))
    with b:
        if avail["events"]:
            fig = go.Figure(go.Bar(x=ctx.month_start, y=ctx.event_count.astype(float),
                                   marker=dict(color=T.GREEN, cornerradius=4),
                                   hovertemplate="%{x|%b}: %{y:,.0f} event links<extra></extra>"))
            fig.update_xaxes(dtick="M1", tickformat="%b")
            show(fig, 240, "Government event links per month")
        else:
            html(T.note("Events: pending", "Appear once event_days is loaded."))

with t_table:
    html(T.section("City-month table", "Every number on this page comes from this table "
                   "(one row per industrial city per month)."))
    table_df = view.drop(columns=["latitude", "longitude"])
    # codes and helper flags stay in the CSV download, not on screen
    on_screen = table_df.drop(columns=[c for c in ("city_code", "region_code", "quarter", "in_both_quarters")
                                       if c in table_df.columns])
    st.dataframe(pretty_table(on_screen), width="stretch", hide_index=True)
    csv = table_df.assign(month_start=table_df.month_start.dt.strftime("%Y-%m"))
    st.download_button("Download CSV", csv.to_csv(index=False).encode("utf-8"), "masar_city_month.csv", "text/csv")

with t_pipe:
    n_files, n_sheets = source_inventory(str(PROJECT_DIR))
    n_models = len(list((PROJECT_DIR / "dbt" / "models").rglob("*.sql")))
    html(T.section("How the data gets here", "Orchestrated by Airflow, provisioned with Terraform, "
                   "run with Docker Compose."))
    html(T.pipeline([
        ("Sources", "Quarterly traffic & event workbooks + Open-Meteo weather API",
         str(n_files) if n_files else "", "Excel files"),
        ("Python ingestion", "Cleans sheets, standardises cities, flags quality issues",
         str(n_sheets) if n_sheets else "", "worksheets read" if n_sheets else ""),
        ("Cloud Storage", "Cleaned files land in the raw bucket", "", ""),
        ("BigQuery raw", "Loaded into the masar_raw dataset", "", ""),
        ("dbt models", "Staging + marts build fct_city_movement_monthly",
         str(n_models) if n_models else "", "models" if n_models else ""),
        ("Streamlit", "This dashboard reads the mart", f"{len(data):,}", "city-months"),
    ]))
    html(T.section("Sources & method", "Where each number comes from"))
    html(T.sources([
        ("Vehicle movements",
         "Industrial-city gate cameras, quarterly files for Q1 and Q2 2026. Entries and exits per city per month. "
         "Missing months are never filled in."),
        ("Weather",
         "Open-Meteo daily weather per industrial city, averaged to city-month. Hot day = max ≥ "
         f"{metrics.HOT_DAY_C:.0f} °C. Context only."),
        ("Government events",
         "Public-sector event approvals. Linked by the city rule: same city → that industrial city; "
         "same region → every industrial city in the region. Context only."),
    ]))
    details = [f"<b>Reading now:</b> {T.esc(SOURCE_SHORT.get(result.source, result.label))}"
               + (f" · <code>{T.esc(data_source.MART_TABLE)}</code>" if live else f" · {T.esc(result.label)}")]
    details += [T.esc(n) for n in result.notes]
    html(T.note("Data source in use", "<br>".join(details), "info" if live else "warn"))

html(T.footer("<b>MASAR</b> · SDA × Le Wagon Data Engineering",
              "Sources &amp; method → <b>Pipeline &amp; sources</b> tab"))